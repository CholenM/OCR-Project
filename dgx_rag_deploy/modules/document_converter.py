"""Document-to-Markdown conversion helpers for RAG ingestion."""

from __future__ import annotations

import csv
import io
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional


SUPPORTED_EXTENSIONS = {
    ".md", ".txt", ".csv",
    ".docx", ".doc", ".dotx", ".odt", ".rtf",
    ".xls", ".xlsx", ".xlsm", ".xlsb", ".xlt",
    ".ppt", ".pptx",
}

SPREADSHEET_EXTENSIONS = {".xls", ".xlsx", ".xlsm", ".xlsb", ".xlt"}
LIBREOFFICE_EXTENSIONS = SUPPORTED_EXTENSIONS - {".md", ".txt", ".csv", ".docx", ".dotx"}


class ConversionError(RuntimeError):
    """Raised when a supported file cannot be converted into useful text."""

    def __init__(self, message: str, status_code: int = 422, warnings: Optional[List[str]] = None):
        super().__init__(message)
        self.status_code = status_code
        self.warnings = warnings or []


@dataclass
class ConversionResult:
    markdown: str
    source_ext: str
    strategy: str
    warnings: List[str] = field(default_factory=list)
    metadata: Dict[str, object] = field(default_factory=dict)


def convert_to_markdown(filename: str, content_bytes: bytes) -> ConversionResult:
    """Convert supported document bytes into Markdown for downstream chunking."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ConversionError(f"Unsupported file format: {filename}", status_code=400)

    if ext in {".md", ".txt"}:
        markdown = content_bytes.decode("utf-8", errors="replace")
        return _validate(ConversionResult(markdown, ext, "plain-text"))

    if ext == ".csv":
        return _validate(ConversionResult(_csv_to_markdown(content_bytes), ext, "csv"))

    if ext in {".docx", ".dotx"}:
        try:
            return _validate(ConversionResult(_docx_to_markdown(content_bytes), ext, "python-docx"))
        except Exception as exc:
            if ext == ".docx":
                return _libreoffice_to_markdown(filename, content_bytes, [f"python-docx failed: {exc}"])
            return _libreoffice_to_markdown(filename, content_bytes, [f"python-docx/dotx failed: {exc}"])

    if ext in SPREADSHEET_EXTENSIONS:
        try:
            return _validate(ConversionResult(_spreadsheet_to_markdown(content_bytes, ext), ext, "pandas"))
        except Exception as exc:
            return _libreoffice_to_markdown(filename, content_bytes, [f"pandas spreadsheet extraction failed: {exc}"])

    return _libreoffice_to_markdown(filename, content_bytes, [])


def _validate(result: ConversionResult) -> ConversionResult:
    text = re.sub(r"\s+", "", result.markdown or "")
    if len(text) < 5:
        raise ConversionError(
            "Document conversion produced no usable text. Try OCR/render fallback for scanned or image-only content.",
            status_code=422,
            warnings=result.warnings,
        )
    return result


def _escape_cell(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r", " ").replace("\n", "<br>").replace("|", "\\|").strip()
    return text


def _rows_to_markdown(rows: List[List[object]]) -> str:
    rows = [[_escape_cell(cell) for cell in row] for row in rows if any(str(cell).strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in rows[1:]:
        out.append("| " + " | ".join(row) + " |")
    return "\n".join(out)


def _csv_to_markdown(content_bytes: bytes) -> str:
    text = content_bytes.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return _rows_to_markdown(rows)


def _docx_to_markdown(content_bytes: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(content_bytes))
    parts: List[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if "heading 1" in style:
            parts.append(f"# {text}")
        elif "heading 2" in style:
            parts.append(f"## {text}")
        elif "heading 3" in style:
            parts.append(f"### {text}")
        elif "heading" in style:
            parts.append(f"#### {text}")
        else:
            parts.append(text)

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md_table = _rows_to_markdown(rows)
        if md_table:
            parts.append(md_table)
    return "\n\n".join(parts)


def _spreadsheet_to_markdown(content_bytes: bytes, ext: str) -> str:
    import pandas as pd

    engine = None
    if ext == ".xlsb":
        engine = "pyxlsb"
    elif ext == ".xls":
        engine = "xlrd"
    data = pd.read_excel(io.BytesIO(content_bytes), sheet_name=None, engine=engine, header=None)
    parts = []
    for sheet_name, frame in data.items():
        frame = frame.dropna(how="all").dropna(axis=1, how="all")
        if frame.empty:
            continue
        rows = frame.fillna("").values.tolist()
        table = _rows_to_markdown(rows)
        if table:
            parts.append(f"## Sheet: {sheet_name}\n\n{table}")
    return "\n\n".join(parts)


def _libreoffice_to_markdown(filename: str, content_bytes: bytes, warnings: List[str]) -> ConversionResult:
    ext = Path(filename).suffix.lower()
    soffice = shutil.which(os.getenv("LIBREOFFICE_BIN", "soffice")) or shutil.which("libreoffice")
    if not soffice:
        raise ConversionError(
            "LibreOffice headless is required for this file type but was not found. Install libreoffice or set LIBREOFFICE_BIN.",
            status_code=422,
            warnings=warnings,
        )

    with tempfile.TemporaryDirectory(prefix="rag_convert_") as tmp:
        tmp_path = Path(tmp)
        src = tmp_path / Path(filename).name
        src.write_bytes(content_bytes)
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        cmd = [
            soffice, "--headless", "--nologo", "--nofirststartwizard",
            "--convert-to", "html", "--outdir", str(out_dir), str(src),
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        if proc.returncode != 0:
            raise ConversionError(
                f"LibreOffice conversion failed for {filename}: {(proc.stderr or proc.stdout).strip()[:300]}",
                status_code=422,
                warnings=warnings,
            )
        html_files = list(out_dir.glob("*.html")) + list(out_dir.glob("*.htm"))
        if not html_files:
            raise ConversionError(
                f"LibreOffice did not produce an HTML output for {filename}.",
                status_code=422,
                warnings=warnings,
            )
        markdown = _html_to_markdown(html_files[0].read_text(encoding="utf-8", errors="replace"))
        return _validate(ConversionResult(markdown, ext, "libreoffice-html", warnings))


def _html_to_markdown(html: str) -> str:
    try:
        import html2text

        h = html2text.HTML2Text()
        h.ignore_links = False
        h.body_width = 0
        return h.handle(html).strip()
    except Exception:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        parts = []
        for table in soup.find_all("table"):
            rows = []
            for tr in table.find_all("tr"):
                rows.append([cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])])
            table_md = _rows_to_markdown(rows)
            if table_md:
                parts.append(table_md)
            table.decompose()
        text = soup.get_text("\n", strip=True)
        if text:
            parts.insert(0, text)
        return "\n\n".join(parts).strip()
